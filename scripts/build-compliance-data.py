"""Builds src/content/compliance/requirements.json from the two published calendars.

  DC PCSB  2026-27 Public Charter LEA Submission Calendar  (authoritative for SY26-27)
  OSSE     2025-26 LEA Requirements Calendar               (SY26-27 edition not yet published;
                                                             dated items are projected one year
                                                             forward and flagged as projected)

  python3 scripts/build-compliance-data.py
"""
import datetime as dt
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
PCSB = ROOT / "data/compliance/pcsb-lea-submission-calendar-sy26-27.xlsx"
OSSE = ROOT / "data/compliance/osse-lea-requirements-sy25-26.xlsx"
OUT = ROOT / "src/content/compliance/requirements.json"

PCSB_URL = "https://dcpcsb.org/annual-compliance-review"
OSSE_URL = "https://osse.dc.gov/publication/2025-26-school-year-lea-requirements-calendar"

MONTHS = {m: i for i, m in enumerate(
    "january february march april may june july august september october november december".split(), 1)}
SEASON_END = {"spring": (5, 31), "summer": (8, 31), "fall": (11, 30), "winter": (2, 28)}


def clean(v):
    if v is None:
        return ""
    return re.sub(r"[ \t\xa0]+", " ", str(v)).strip()


def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")[:60]


def parse_pcsb_due(v):
    """-> (iso date or None, text label, precision)"""
    if isinstance(v, dt.datetime):
        return v.date().isoformat(), "", "day"
    s = clean(v)
    m = re.match(r"(\d{1,2})/(\d{1,2})/(\d{4})", s)
    if m:
        return dt.date(int(m[3]), int(m[1]), int(m[2])).isoformat(), "", "day"
    m = re.match(r"(spring|summer|fall|winter) (\d{4})", s, re.I)
    if m:
        mo, d = SEASON_END[m[1].lower()]
        return dt.date(int(m[2]), mo, d).isoformat(), s, "season"
    m = re.match(r"([A-Za-z]+) (\d{1,2}) - ([A-Za-z]+) (\d{1,2})", s)
    if m and m[3].lower() in MONTHS:
        mo = MONTHS[m[3].lower()]
        yr = 2026 if mo >= 7 else 2027
        return dt.date(yr, mo, int(m[4])).isoformat(), f"{m[1]} {m[2]} – {m[3]} {m[4]}", "window"
    return None, s.split("\n")[0], "none"


def pcsb_applies(who, title):
    w = who.lower()
    if "high school" in w or "not participating in msdc" in w:
        return "no", "High schools only" if "high" in w else "Only for schools outside My School DC"
    if "amendment" in w or "waiver" in w:
        return "if", "Only if you request an amendment or waiver"
    if "identified" in w:
        return "if", "Only if DC PCSB notifies you"
    if "relocating" in w or "new campus" in w:
        return "yes", "You are moving buildings for 2026-27, so this applies"
    return "yes", ""


def owner_for(platform, dashboard, title):
    t = title.lower()
    if "financial" in dashboard.lower() or re.search(r"audit|budget|financ|990|procurement|contract", t):
        return "Finance"
    if re.search(r"enrollment|roster|attendance|discipline|calendar|entity|data", t):
        return "Data & enrollment"
    if re.search(r"board|charter|annual report|policy|title ix|accreditation", t):
        return "Head of school"
    return "Operations"


def link_of(cell):
    try:
        return cell.hyperlink.target if cell.hyperlink else ""
    except Exception:
        return ""


def load_pcsb():
    ws = openpyxl.load_workbook(PCSB).worksheets[0]
    rows = [r for r in ws.iter_rows() if any(c.value is not None for c in r)]
    head = [clean(c.value).split(" (")[0].split("\n")[0] for c in rows[1]]
    col = {h: i for i, h in enumerate(head)}
    out = []
    for r in rows[2:]:
        g = lambda h: clean(r[col[h]].value)
        title = g("Submission Title")
        if not title:
            continue
        due, due_label, precision = parse_pcsb_due(r[col["Due Date"]].value)
        platform = g("Collection Platform")
        if platform.startswith("The Hub"):
            platform = "The Hub"
        elif "email" in platform.lower():
            platform = "Email to DC PCSB"
        applies, why = pcsb_applies(g("Which Schools Are Required to Submit?"), title)
        contact = g("DC PCSB Contact").replace("\n", " · ")
        rule_cell = r[col["Business Rule"]]
        dash = g("Dashboard Location for Tasks Collected in The Hub") if "Dashboard Location for Tasks Collected in The Hub" in col else clean(r[10].value)
        out.append({
            "id": "pcsb-" + slug(title),
            "agency": "DC PCSB",
            "kind": g("Type") or "Submission",
            "title": title,
            "due": due,
            "dueLabel": due_label,
            "precision": precision,
            "projected": False,
            "purpose": g("Purpose"),
            "guide": clean(rule_cell.value),
            "guideUrl": link_of(rule_cell),
            "whoMustSubmit": g("Which Schools Are Required to Submit?"),
            "applies": applies,
            "appliesWhy": why,
            "platform": platform,
            "contact": contact,
            "alsoOsse": g("Requested by OSSE").lower().startswith("yes"),
            "owner": owner_for(platform, dash or "", title),
            "source": PCSB_URL,
        })
    return out


OSSE_SKIP = re.compile(
    r"high school|DCPS ONLY|college|grade(s)? 9|9th|10th|11th|12th|graduat|career ready|\bcte\b|perkins|adult and family|"
    r"\bafe\b|laces|dual enrollment|postsecondary|psat|\bsat\b|advanced placement|diploma|transcript",
    re.I)


def osse_owner(role):
    """OSSE names ~20 LEA roles; a single-campus school has four or five people behind them."""
    r = role.lower()
    if "head of school" in r:
        return "Head of school"
    if re.search(r"finance|grant", r):
        return "Finance"
    if re.search(r"special ed|idea|504", r):
        return "Special education"
    if re.search(r"data|enrollment|assessment|access for ells|adt|sis", r):
        return "Data & enrollment"
    return "Operations"


def load_osse():
    ws = openpyxl.load_workbook(OSSE, read_only=True)["Normalized"]
    rows = list(ws.iter_rows(values_only=True))
    head = [clean(h) for h in rows[0]]
    col = {h: i for i, h in enumerate(head)}
    out = []
    for r in rows[1:]:
        g = lambda h: r[col[h]]
        title = clean(g("Item"))
        if not title:
            continue
        desc = clean(g("Description"))
        kind = clean(g("DeadlineType"))
        end = g("End_date") or (g("SortEnd") if kind == "date" else None)
        due, label, precision, projected = None, clean(g("PublicDisplay")), "none", False
        if isinstance(end, dt.datetime) and kind == "date":
            # 52 weeks, not a calendar year: OSSE sets deadlines on weekdays, and the same
            # date a year later lands a day later in the week (Fri 10 Oct 2025 -> Sat 10 Oct 2026).
            due = (end.date() + dt.timedelta(days=364)).isoformat()
            label, precision, projected = "", "day", True
        applies, why = "yes", ""
        if OSSE_SKIP.search(title + " " + desc):
            applies, why = "no", "Not for a charter elementary school (DCPS, high school, career or adult programs)"
        elif kind in ("trigger_based", "as_needed"):
            applies, why = "if", "Only when the triggering event happens"
        platform = clean(g("Submission Platform")) or "See guidance"
        platform = {
            "TBD - program to supply": "Platform not yet announced",
            "Email to Program": "Email to the OSSE program office",
            "Varies": "Varies by program",
        }.get(platform, platform)
        out.append({
            "id": "osse-" + slug(title),
            "agency": "OSSE",
            "kind": "Submission",
            "title": re.sub(r"\s*\(.*?\)\s*$", "", title) if len(title) > 70 else title,
            "due": due,
            "dueLabel": label.replace("SY25–26", "the school year").replace("SY25-26", "the school year"),
            "precision": precision,
            "projected": projected,
            "purpose": desc,
            "guide": "",
            "guideUrl": clean(g("Related Links")) if clean(g("Related Links")).startswith("http") else "",
            "whoMustSubmit": clean(g("LEA Role")),
            "applies": applies,
            "appliesWhy": why,
            "platform": platform,
            "contact": " · ".join(x for x in [clean(g("OSSE POC")), clean(g("OSSE POC Email"))] if x),
            "alsoOsse": True,
            "owner": osse_owner(clean(g("LEA Role"))),
            "source": OSSE_URL,
        })
    return out


def main():
    # OSSE lists one requirement once per submission channel (the dyslexia screening
    # appears four times on one date). Keep one row and name the other channels on it.
    merged = {}
    for it in load_osse():
        k = (it["title"], it["due"])
        if k in merged:
            m = merged[k]
            if it["platform"] not in m["platform"]:
                m["alsoVia"] = sorted(set(m.get("alsoVia", []) + [it["platform"]]))
        else:
            merged[k] = it
    items = load_pcsb() + list(merged.values())
    seen = {}
    for it in items:
        n = seen.get(it["id"], 0)
        seen[it["id"]] = n + 1
        if n:
            it["id"] += f"-{n + 1}"
    items.sort(key=lambda x: (x["due"] or "9999", x["agency"], x["title"]))
    OUT.write_text(json.dumps({
        "generated": dt.date.today().isoformat(),
        "sources": [
            {"label": "DC PCSB 2026-27 Public Charter LEA Submission Calendar", "url": PCSB_URL},
            {"label": "OSSE 2025-26 LEA Requirements Calendar (March 2026 edition)", "url": OSSE_URL},
        ],
        "items": items,
    }, indent=1, ensure_ascii=False))
    from collections import Counter
    print(len(items), "items", Counter((i["agency"], i["applies"]) for i in items))
    print("platforms", Counter(i["platform"] for i in items if i["applies"] != "no").most_common(40))


main()
